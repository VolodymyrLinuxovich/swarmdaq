#!/usr/bin/env python3
"""Generate a clean SwarmDAQ pitch deck with safe text sizing (no clipping)."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

# ── palette ──────────────────────────────────────────────────────────────────
BG = RGBColor(10, 15, 13)
PANEL = RGBColor(18, 26, 23)
PANEL2 = RGBColor(24, 34, 30)
STROKE = RGBColor(40, 58, 52)
TEXT = RGBColor(220, 235, 230)
MUTED = RGBColor(130, 155, 145)
DIM = RGBColor(90, 115, 105)
GREEN = RGBColor(0, 217, 122)
BLUE = RGBColor(59, 158, 255)
AMBER = RGBColor(255, 176, 32)
RED = RGBColor(255, 90, 100)
PURPLE = RGBColor(168, 108, 255)

W = Inches(13.333)
H = Inches(7.5)
ML = Inches(0.8)
MR = Inches(0.8)
MT = Inches(0.5)
FOOTER_Y = Inches(6.95)
CONTENT_W = W - ML - MR
FONT = "Arial"


def _fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def _line(shape, color, width=1):
    shape.line.color.rgb = color
    shape.line.width = Pt(width)


def _no_line(shape):
    shape.line.fill.background()


def _rect(slide, left, top, width, height, fill=PANEL, stroke=STROKE):
    shape = slide.shapes.add_shape(1, left, top, width, height)
    _fill(shape, fill)
    if stroke:
        _line(shape, stroke)
    else:
        _no_line(shape)
    return shape


def _lines_for_text(text, width_in, size_pt):
    """Rough line-count estimate for wrapped text."""
    cpl = max(12, int(width_in * 72 / (size_pt * 0.52)))
    total = 0
    for part in text.split("\n"):
        total += max(1, (len(part) + cpl - 1) // cpl)
    return total


def _h_in(lines, size_pt, leading=1.3, pad_in=0.2):
    # 1.4x buffer — PowerPoint/Keynote need more room than naive line math
    return Inches((lines * (size_pt / 72) * leading + pad_in) * 1.4)


def _textbox(slide, left, top, width, text, size=14, color=TEXT, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, leading=1.3,
             height=None, auto_fit=False):
    width_in = width / 914400
    if height is None and not auto_fit:
        height = _h_in(_lines_for_text(text, width_in, size), size, leading)
    if auto_fit:
        height = Inches(0.2)

    # Never clip ascenders/descenders on single-line text
    min_h = Inches(size / 72 * 1.75 + 0.14)
    if not auto_fit and height < min_h:
        height = min_h

    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Pt(8)
    tf.margin_right = Pt(8)
    tf.margin_top = Pt(6)
    tf.margin_bottom = Pt(6)

    if auto_fit:
        tf.auto_size = MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
    else:
        tf.auto_size = MSO_AUTO_SIZE.NONE

    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = FONT
    p.alignment = align
    p.line_spacing = leading
    return box


def _bullets(slide, left, top, width, items, size=14, color=TEXT, leading=1.35):
    lines = 0
    width_in = width / 914400
    for item in items:
        lines += _lines_for_text(f"• {item}", width_in, size)
    height = _h_in(lines, size, leading, pad_in=0.08)

    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = Pt(8)
    tf.margin_right = Pt(8)
    tf.margin_top = Pt(6)
    tf.margin_bottom = Pt(6)

    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"• {item}"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.font.name = FONT
        p.line_spacing = leading
        p.space_before = Pt(6) if i else Pt(0)
    return box


def _bg(slide):
    bg = slide.shapes.add_shape(1, 0, 0, W, H)
    _fill(bg, BG)
    _no_line(bg)
    bar = slide.shapes.add_shape(1, 0, 0, W, Inches(0.035))
    _fill(bar, GREEN)
    _no_line(bar)


def _footer(slide, page, source=""):
    _textbox(slide, ML, FOOTER_Y, Inches(9.5), source, size=9, color=DIM, height=Inches(0.35))
    _textbox(slide, Inches(12.1), FOOTER_Y, Inches(0.6), f"{page:02d}",
             size=9, color=DIM, align=PP_ALIGN.RIGHT, height=Inches(0.35))


def _header(slide, kicker, title, subtitle=None):
    y = MT
    if kicker:
        _textbox(slide, ML, y, Inches(3), kicker.upper(), size=10, color=GREEN, bold=True, height=Inches(0.22))
        y = Inches(0.72)

    _textbox(slide, ML, y, CONTENT_W, title, size=26, color=TEXT, bold=True,
             leading=1.2, height=_h_in(_lines_for_text(title, CONTENT_W / 914400, 26), 26, 1.2, 0.1))

    if subtitle:
        sub_y = y + Inches(0.95)
        _textbox(slide, ML, sub_y, CONTENT_W, subtitle, size=14, color=MUTED,
                 leading=1.35, height=_h_in(_lines_for_text(subtitle, CONTENT_W / 914400, 14), 14, 1.35, 0.1))
        return sub_y + Inches(0.55)
    return y + Inches(0.85)


def _card(slide, left, top, width, title, body, accent=GREEN, title_size=15, body_size=13):
    inner_w = width - Inches(0.4)
    inner_in = inner_w / 914400
    title_h = Inches(0.32)
    body_h = _h_in(_lines_for_text(body, inner_in, body_size), body_size, 1.4, 0.15)
    height = title_h + body_h + Inches(0.55)
    _rect(slide, left, top, width, height, fill=PANEL2, stroke=accent)
    _textbox(slide, left + Inches(0.2), top + Inches(0.15), inner_w, title,
             size=title_size, color=accent, bold=True, height=title_h)
    _textbox(slide, left + Inches(0.2), top + Inches(0.48), inner_w, body,
             size=body_size, color=TEXT, leading=1.4, height=body_h)
    return height


# ── slides ───────────────────────────────────────────────────────────────────

def slide_cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(slide)

    _textbox(slide, ML, Inches(1.5), CONTENT_W, "LIVE TECHNICAL DEMO",
             size=11, color=GREEN, bold=True, height=Inches(0.25))
    _textbox(slide, ML, Inches(1.9), Inches(8), "SwarmDAQ",
             size=42, color=GREEN, bold=True, height=Inches(0.75))
    _textbox(slide, ML, Inches(2.65), Inches(8), "The performance exchange for AI agents",
             size=22, color=TEXT, bold=True, height=Inches(0.45))
    _textbox(slide, ML, Inches(3.2), Inches(8),
             "Agents bid on tasks. A MarketMaker routes with math. "
             "Every run is traced, evaluated, and stored — so the next swarm improves.",
             size=14, color=MUTED, leading=1.4, height=Inches(0.9))

    stats = [
        ("9", "specialized agents\nwith persistent state", GREEN),
        ("7+", "routing signals beyond\nprompt selection", BLUE),
        ("4", "run demo proving\nthe feedback loop", AMBER),
    ]
    card_w = Inches(3.7)
    y = Inches(4.85)
    for i, (val, label, color) in enumerate(stats):
        x = ML + i * (card_w + Inches(0.25))
        _rect(slide, x, y, card_w, Inches(1.55), fill=PANEL2, stroke=STROKE)
        _textbox(slide, x + Inches(0.25), y + Inches(0.2), card_w - Inches(0.5), val,
                 size=34, color=color, bold=True, height=Inches(0.55))
        _textbox(slide, x + Inches(0.25), y + Inches(0.8), card_w - Inches(0.5), label,
                 size=13, color=MUTED, leading=1.3, height=Inches(0.6))

    _footer(slide, 1, "Source: SwarmDAQ repo, paper, and live demo.")


def slide_problem(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(slide)
    content_y = _header(slide, "Problem",
                        "Agent teams fail when trust is invisible.",
                        "Most orchestration behaves like a prompt chain — assign work, hope for the best, lose the evidence.")

    items = [
        "Static routing with no performance memory",
        "Weak outputs go undetected — unsupported claims, bad fit",
        "Same agent gets reused even after failure",
        "No answer to: who performed, who hallucinated, who to trust next?",
    ]
    _bullets(slide, ML, content_y, Inches(6.5), items)

    _card(slide, ML, Inches(5.0), Inches(6.5),
          "The missing layer",
          "A transparent exchange where agents earn or lose trust from actual performance — not static role labels.",
          GREEN)
    _card(slide, Inches(7.7), Inches(5.0), Inches(4.85),
          "Demo failure (Run 1)",
          "Unsupported claim detected. ResearchAgent penalized. SourceVerifierAgent promoted.",
          RED)
    _footer(slide, 2, "Source: README problem framing.")


def slide_solution(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(slide)
    content_y = _header(slide, "Solution",
                        "The product is a market loop, not another agent.",
                        "Bid → select → execute → trace → evaluate → update.")

    steps = [
        "Mission — user intent enters the system",
        "Task graph — PlannerAgent decomposes work",
        "Agent bids — confidence, cost, latency, utility",
        "MarketMaker — weighted selection picks the swarm",
        "Execution — Gemini or deterministic fallback",
        "Evaluation — six quality dimensions scored",
        "Memory — Beta, Elo, graph trust stored in Redis",
    ]
    _bullets(slide, ML, content_y, Inches(6.8), steps, size=14)

    _card(slide, Inches(8.0), content_y, Inches(4.55),
          "After every run",
          "Reputation updates.\nUncertainty tracked.\nElo ratings shift.\nCollaboration graph grows.\nMission replay saved.",
          GREEN, body_size=14)
    _footer(slide, 3, "Source: orchestrator.ts, memory.ts, trace.ts.")


def slide_architecture(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(slide)
    content_y = _header(slide, "Architecture",
                        "End-to-end: UI, routing, tracing, and memory.")

    rows = [
        ("Frontend", "Next.js — landing, demo, leaderboard, benchmark"),
        ("API", "/api/mission, /agents, /traces, /history, /market-feed"),
        ("Orchestrator", "Planning, bidding, selection, execution, reputation"),
        ("Execution", "Gemini 2.5 Flash + deterministic fallback"),
        ("Evaluator", "Quality, factuality, usefulness, actionability, collaboration"),
        ("Tracing", "W&B Weave + local trace event fallback"),
        ("Memory", "Upstash Redis or in-memory reputation state"),
    ]

    row_h = Inches(0.56)
    for i, (label, value) in enumerate(rows):
        y = content_y + row_h * i
        bg = PANEL if i % 2 == 0 else PANEL2
        _rect(slide, ML, y, CONTENT_W, row_h, fill=bg, stroke=STROKE)
        _textbox(slide, ML + Inches(0.18), y + Inches(0.1), Inches(1.8), label,
                 size=13, color=GREEN, bold=True, height=Inches(0.38))
        _textbox(slide, ML + Inches(2.1), y + Inches(0.1), CONTENT_W - Inches(2.3), value,
                 size=13, color=TEXT, height=Inches(0.38))

    note_y = content_y + row_h * len(rows) + Inches(0.2)
    _card(slide, ML, note_y, CONTENT_W,
          "Key design choice",
          "Every production dependency has a fallback path — the demo works without API keys.",
          GREEN)
    _footer(slide, 4, "Source: Next.js app routes and lib/.")


def slide_routing(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(slide)
    content_y = _header(slide, "Routing math",
                        "MarketMaker routes with a score, not vibes.")

    _rect(slide, ML, content_y, Inches(6.2), Inches(1.65), fill=PANEL2, stroke=GREEN)
    _textbox(slide, ML + Inches(0.2), content_y + Inches(0.15), Inches(5.8),
             "finalScore =", size=16, color=GREEN, bold=True, height=Inches(0.3))
    _textbox(slide, ML + Inches(0.2), content_y + Inches(0.45), Inches(5.8),
             "0.20·skill + 0.18·reputation + 0.16·UCB + 0.14·bid\n"
             "+ 0.12·Elo + 0.08·graphTrust + 0.07·collaboration\n"
             "+ 0.05·factuality − 0.05·cost − 0.05·latency − 0.10·uncertainty",
             size=13, color=TEXT, leading=1.4, height=Inches(1.1))

    weights = [
        ("+0.20", "Skill match"),
        ("+0.18", "Bayesian reputation"),
        ("+0.16", "UCB1 exploration"),
        ("+0.14", "Bid utility"),
        ("+0.12", "Elo rating"),
        ("+0.08", "Graph trust"),
        ("−0.20", "Cost, latency, uncertainty penalties"),
    ]
    wy = content_y
    for i, (w, name) in enumerate(weights):
        y = wy + Inches(0.5) * i
        bg = PANEL if i % 2 == 0 else PANEL2
        _rect(slide, Inches(7.5), y, Inches(5.05), Inches(0.46), fill=bg, stroke=STROKE)
        color = RED if w.startswith("−") else GREEN
        _textbox(slide, Inches(7.65), y + Inches(0.1), Inches(0.7), w,
                 size=12, color=color, bold=True, height=Inches(0.28))
        _textbox(slide, Inches(8.4), y + Inches(0.1), Inches(3.8), name,
                 size=13, color=TEXT, height=Inches(0.28))

    _footer(slide, 5, "Source: computeMarketMakerScore() in orchestrator.ts.")


def slide_agents(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(slide)
    content_y = _header(slide, "Agent market",
                        "Agents carry market state like real participants.",
                        "Price, latency, factuality, Beta trust, Elo, runs, wins, losses.")

    agents = [
        ("PlannerAgent", "Mission decomposition"),
        ("ResearchAgent", "Market analysis"),
        ("SourceVerifierAgent", "Claim validation"),
        ("BuilderAgent", "Product synthesis"),
        ("PitchAgent", "Narrative and copy"),
        ("SkepticAgent", "Risk analysis"),
        ("EvaluatorAgent", "Quality scoring"),
        ("MarketMakerAgent", "Routing and auctions"),
        ("ReputationAgent", "Memory updates"),
    ]

    row_h = Inches(0.42)
    _rect(slide, ML, content_y, CONTENT_W, row_h, fill=PANEL2, stroke=STROKE)
    _textbox(slide, ML + Inches(0.15), content_y + Inches(0.08), Inches(3), "Agent",
             size=11, color=DIM, bold=True, height=Inches(0.28))
    _textbox(slide, ML + Inches(3.2), content_y + Inches(0.08), Inches(4), "Role",
             size=11, color=DIM, bold=True, height=Inches(0.28))

    for i, (name, role) in enumerate(agents):
        y = content_y + row_h * (i + 1)
        bg = PANEL if i % 2 == 0 else PANEL2
        _rect(slide, ML, y, CONTENT_W, row_h, fill=bg, stroke=STROKE)
        _textbox(slide, ML + Inches(0.15), y + Inches(0.08), Inches(3), name,
                 size=12, color=TEXT, bold=True, height=Inches(0.28))
        _textbox(slide, ML + Inches(3.2), y + Inches(0.08), Inches(6), role,
                 size=12, color=MUTED, height=Inches(0.28))

    _card(slide, ML, Inches(6.05), CONTENT_W,
          "Routing is inspectable",
          "Every selected agent has history, price, risk profile, and an explanation.",
          BLUE)
    _footer(slide, 6, "Source: agents.ts and types.ts.")


def slide_demo(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(slide)
    content_y = _header(slide, "Demo proof",
                        "Learning, regression, and recovery.",
                        "Internal demo values (74 → 91 → 85 → 96). Evaluation changes memory; memory changes routing.")

    runs = [("Run 1", 74, "failure", RED), ("Run 2", 91, "learned", GREEN),
            ("Run 3", 85, "over-rotate", AMBER), ("Run 4", 96, "recovery", BLUE)]

    chart_x, chart_y = ML, content_y + Inches(0.1)
    chart_w, chart_h = Inches(7.0), Inches(3.2)
    _rect(slide, chart_x, chart_y, chart_w, chart_h, fill=PANEL, stroke=STROKE)

    bar_w = Inches(1.2)
    gap = Inches(0.35)
    plot_h = Inches(2.0)
    base = chart_y + Inches(2.55)

    for i, (label, score, note, color) in enumerate(runs):
        x = chart_x + Inches(0.5) + i * (bar_w + gap)
        h = plot_h * (score / 100)
        _rect(slide, x, base - h, bar_w, h, fill=color, stroke=None)
        _textbox(slide, x, base - h - Inches(0.4), bar_w, str(score),
                 size=20, color=color, bold=True, align=PP_ALIGN.CENTER, height=Inches(0.35))
        _textbox(slide, x, base + Inches(0.08), bar_w, label,
                 size=12, color=TEXT, bold=True, align=PP_ALIGN.CENTER, height=Inches(0.25))
        _textbox(slide, x - Inches(0.1), base + Inches(0.32), bar_w + Inches(0.2), note,
                 size=10, color=MUTED, align=PP_ALIGN.CENTER, height=Inches(0.3))

    notes = [
        "Run 2: SourceVerifierAgent promoted after factuality failure",
        "Run 3: Over-optimization — skepticism hurts narrative",
        "Run 4: Swarm rebalances, best internal score",
    ]
    _bullets(slide, Inches(8.3), content_y + Inches(0.3), Inches(4.3), notes, size=13)
    _footer(slide, 7, "Source: internal demo values in orchestrator.ts.")


def slide_dashboard(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(slide)
    content_y = _header(slide, "Product",
                        "The dashboard makes the market inspectable.")

    features = [
        "Mission input with live agent bidding",
        "MarketMaker score breakdown per task",
        "Live deliberation: objections and revisions",
        "W&B Weave trace timeline",
        "Six-dimension evaluation scorecards",
        "UCB, Bayesian, and Elo math panels",
        "Leaderboard with BUY / HOLD / SELL signals",
        "Agent Studio for custom agent deployment",
    ]
    _bullets(slide, ML, content_y, Inches(6.5), features, size=13)

    _card(slide, Inches(8.0), content_y, Inches(4.55),
          "What users see",
          "Why the swarm changed.\nWhich agent won each task.\nHow reputation shifted.\nFull mission replay.",
          GREEN)
    _footer(slide, 8, "Source: demo page and dashboard components.")


def slide_wedge(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(slide)
    content_y = _header(slide, "Wedge",
                        "Trust infrastructure for the agent economy.")

    cards = [
        ("Why now", "Agent supply is exploding. Buyers lack transparent performance, routing, and failure memory.", AMBER),
        ("Why this wins", "Every run becomes market evidence: bids, traces, evals, reputation deltas.", GREEN),
        ("What expands", "External benchmarks, agent catalogs, real budgets, audited production traces.", BLUE),
    ]
    card_w = Inches(3.75)
    for i, (title, body, accent) in enumerate(cards):
        _card(slide, ML + i * (card_w + Inches(0.28)), content_y, card_w,
              title, body, accent, body_size=13)

    _card(slide, ML, Inches(5.35), CONTENT_W,
          "Positioning",
          "Not a chatbot. Not a generic framework. The exchange layer that decides which agents are worth hiring.",
          GREEN)
    _footer(slide, 9, "Source: README and paper.")


def slide_close(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(slide)

    _textbox(slide, ML, Inches(1.7), CONTENT_W, "Make agent routing auditable.",
             size=36, color=TEXT, bold=True, height=Inches(0.9))
    _textbox(slide, ML, Inches(2.55), Inches(8),
             "Every agent gets a market price, performance history, trace, and reason to improve.",
             size=18, color=GREEN, bold=True, leading=1.35, height=Inches(0.75))

    links = [
        ("Live demo", "swarmdaq.vercel.app", GREEN),
        ("Repository", "github.com/VolodymyrLinuxovich/swarmdaq", BLUE),
        ("Paper", "paper/main.tex + Actions PDF", AMBER),
    ]
    card_w = Inches(3.75)
    y = Inches(3.85)
    for i, (title, url, accent) in enumerate(links):
        x = ML + i * (card_w + Inches(0.28))
        _rect(slide, x, y, card_w, Inches(1.3), fill=PANEL2, stroke=accent)
        _textbox(slide, x + Inches(0.22), y + Inches(0.18), card_w - Inches(0.44), title,
                 size=15, color=accent, bold=True, height=Inches(0.3))
        _textbox(slide, x + Inches(0.22), y + Inches(0.52), card_w - Inches(0.44), url,
                 size=12, color=TEXT, height=Inches(0.55))

    _textbox(slide, ML, Inches(5.5), CONTENT_W,
             "Next: benchmark suite · agent catalog · budget-aware clearing · adversarial bid tests",
             size=13, color=MUTED, leading=1.35, height=Inches(0.55))
    _footer(slide, 10, "swarmdaq.vercel.app")


def build(output_path: Path):
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    for fn in (slide_cover, slide_problem, slide_solution, slide_architecture,
               slide_routing, slide_agents, slide_demo, slide_dashboard,
               slide_wedge, slide_close):
        fn(prs)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
    print(f"Saved {output_path} ({output_path.stat().st_size:,} bytes, {len(prs.slides)} slides)")


if __name__ == "__main__":
    build(Path(__file__).resolve().parent / "swarmdaq-pitchdeck.pptx")
